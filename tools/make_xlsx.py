import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

rows = list(csv.DictReader(open('apartments_full.csv')))
for r in rows:
    r['building'] = int(r['building']); r['apt'] = int(r['apt'])
    r['floor'] = int(r['floor']); r['rooms'] = float(r['rooms']) if r['rooms'] else None
    r['area'] = float(r['area']) if r['area'] else None
    r['balcony_m2'] = float(r['balcony_m2']) if r['balcony_m2'] else 0
    r['balconies'] = int(r['balconies']) if r['balconies'] else 0
    r['yard'] = 'כן' if r['type'] in ('1G', '3G') else ''
    east = 'מזרח' in r['direction']
    r['park'] = 'כן' if (east and r['building'] in (2, 3)) else ('חלקי' if east else '')

prog = sorted([r for r in rows if r['program'] == 'Y'],
              key=lambda r: (r['building'], r['apt']))
assert len(prog) == 84, len(prog)

wb = Workbook()
F = 'Arial'
hdr_fill = PatternFill('solid', fgColor='1F4E78')
hdr_font = Font(name=F, bold=True, color='FFFFFF', size=10)
in_font = Font(name=F, color='0000FF', size=10)
yellow = PatternFill('solid', fgColor='FFFF00')
thin = Border(*[Side(style='thin', color='BFBFBF')]*4)
NIS = '₪#,##0'

# ---------------- Assumptions ----------------
ws = wb.active; ws.title = 'הנחות'
ws.sheet_view.rightToLeft = True
ws['A1'] = 'הנחות וקבועים — עדכנו את התאים הצהובים בלבד'
ws['A1'].font = Font(name=F, bold=True, size=13)
params = [
    ('מחיר למ"ר לפני מע"מ (לפי המכרז)', 15022, NIS, 'נתון מהמשתמש — מכרז מחיר מטרה'),
    ('מקדם מרפסת לחישוב המחיר', 0.3, '0.0%', 'הנחת עבודה מקובלת במכרזי מחיר מטרה/למשתכן — לאמת מול נספח התמורה בחוזה!'),
    ('מע"מ', 0.18, '0.0%', 'שיעור מע"מ נכון ל-2026 — לעדכן אם ישתנה עד החתימה'),
    ('שיעור הנחה', 0.20, '0.0%', 'הנמוך מבין 20% מהמחיר או תקרת ההנחה'),
    ('תקרת הנחה', 300000, NIS, 'נתון מהמשתמש'),
    ('מחיר שוק משוער למ"ר (כולל מע"מ)', 24000, NIS, 'הערכת מקום ראשונית לדירה חדשה בנוף צמרת/פ"ת — חובה לעדכן לפי עסקאות אמת/שמאות'),
    ('מקדם מרפסת בשווי שוק', 0.5, '0.0%', 'הנחת עבודה לשווי מרפסת בשוק החופשי'),
]
ws.append([])
ws.append(['פרמטר', 'ערך', 'הערות'])
for c in ws[3]: c.fill = hdr_fill; c.font = hdr_font
for name, val, fmt, note in params:
    ws.append([name, val, note])
    r = ws.max_row
    ws.cell(r, 1).font = Font(name=F, size=10)
    v = ws.cell(r, 2); v.font = in_font; v.fill = yellow; v.number_format = fmt
    ws.cell(r, 3).font = Font(name=F, size=9, italic=True)
notes = [
    '',
    'הערות חשובות:',
    '1. על המחיר החוזי תתווסף הצמדה למדד תשומות הבנייה (לא מגולמת כאן) — ההצמדה חלה בדרך כלל על יתרת התשלומים.',
    '2. הנתונים חולצו אוטומטית מקובץ ה-DWFX של תכניות המכר (גיליון R18367-X-XX-DP-A-001, Revit 2024, פרויקט "אדרת בצמרת - סירקין", מגרש 201).',
    '3. שטחי מרפסות נלקחו מתוויות "מרפסת שמש" בתכניות, ואומתו פר טיפוס דירה על פני כל הקומות.',
    '4. "פונה לפארק": הפארק הציבורי גובל במגרש ממזרח. ההנחה: בניינים 2-3 בצד המזרחי של המגרש — כיוון מזרח בהם פתוח לפארק; בבניין 1 מזרח מוסתר חלקית. לאמת מול תשריט הפיתוח!',
    '5. בקומה טיפוסית 2-3 מוצגת תכנית אחת לשתי הקומות; מספרי הדירות בקומה 3 = קומה 2 + 5 (אומת מול טבלאות התכנית).',
    '6. דירות 3G (בניין 3, דירות 1-2) כוללות חצר פרטית ששטחה אינו מסומן בתכנית.',
    '7. שדה "תקרת מחיר מטרה" = הכיתוב "שייכות מחיר מטרה עד 80/110 מ"ר" המופיע בטבלת כל דירה בתכנית.',
]
for n in notes:
    ws.append([n]); ws.cell(ws.max_row, 1).font = Font(name=F, size=9)
ws.column_dimensions['A'].width = 38
ws.column_dimensions['B'].width = 14
ws.column_dimensions['C'].width = 95

# ---------------- main sheet ----------------
HDRS = ['בניין', 'דירה', 'קומה', 'טיפוס', 'חדרים', 'שטח דירה (מ"ר)',
        'מרפסת שמש (מ"ר)', 'מס\' מרפסות', 'חצר', 'כיווני אוויר', 'פונה לפארק',
        'תקרת מחיר מטרה (מ"ר)', 'שטח לתמחור (מ"ר)', 'מחיר לפני מע"מ',
        'מחיר כולל מע"מ', 'הנחה', 'מחיר נטו לתשלום', 'שווי שוק משוער',
        'רווח גולמי משוער', 'דירוג רווח']

def fill_sheet(ws, data, with_prices=True):
    ws.sheet_view.rightToLeft = True
    hdrs = HDRS if with_prices else HDRS[:12] + ['מחיר מטרה']
    ws.append(hdrs)
    for c in ws[1]:
        c.fill = hdr_fill; c.font = hdr_font
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    n = len(data)
    for i, r in enumerate(data):
        row = i + 2
        base = [r['building'], r['apt'], r['floor'], r['type'], r['rooms'],
                r['area'], r['balcony_m2'] or None, r['balconies'] or None,
                r['yard'], r['direction'], r['park'], int(r['cap']) if r['cap'] else None]
        if not with_prices:
            ws.append(base + ['כן' if r['program'] == 'Y' else ''])
        else:
            M = f"=F{row}+הנחות!$B$5*G{row}"
            price_ex = f"=ROUND(M{row}*הנחות!$B$4,0)"
            price_inc = f"=ROUND(N{row}*(1+הנחות!$B$6),0)"
            disc = f"=MIN(הנחות!$B$7*O{row},הנחות!$B$8)"
            net = f"=O{row}-P{row}"
            mkt = f"=ROUND(הנחות!$B$9*(F{row}+הנחות!$B$10*G{row}),0)"
            profit = f"=R{row}-Q{row}"
            rank = f"=RANK(S{row},S$2:S${n+1})"
            ws.append(base + [M, price_ex, price_inc, disc, net, mkt, profit, rank])
        for j in range(1, len(hdrs) + 1):
            c = ws.cell(row, j); c.font = Font(name=F, size=10); c.border = thin
        ws.cell(row, 13).number_format = '0.0'
        for j in range(14, 20):
            ws.cell(row, j).number_format = NIS
    widths = [7, 7, 7, 8, 8, 12, 12, 9, 7, 14, 10, 12, 11, 13, 13, 12, 13, 13, 13, 9]
    for j, w in enumerate(widths[:len(hdrs)], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(hdrs))}{len(data)+1}'

# assumptions param rows start at row 4: B4=price/m2, B5=balcony coef, B6=VAT,
# B7=discount rate, B8=discount cap, B9=market price, B10=market balcony coef
ws2 = wb.create_sheet('דירות מחיר מטרה')
fill_sheet(ws2, prog, with_prices=True)
ws3 = wb.create_sheet('כל הדירות בפרויקט')
fill_sheet(ws3, sorted(rows, key=lambda r: (r['building'], r['apt'])), with_prices=False)

wb.save('adereet_apartments.xlsx')
print('saved', len(prog), 'program apts,', len(rows), 'total')
